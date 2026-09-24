from docx import Document


def add_markdown_content(document: Document, content: str):
    lines = content.splitlines()

    i = 0

    while i < len(lines):

        line = lines[i].strip()

        # 空行
        if not line:
            i += 1
            continue

        # Markdown 表格
        if (
            line.startswith("|")
            and i + 1 < len(lines)
            and "|" in lines[i + 1]
        ):
            table_lines = []

            while (
                i < len(lines)
                and lines[i].strip().startswith("|")
            ):
                table_lines.append(lines[i])
                i += 1

            add_markdown_table(
                document,
                table_lines
            )

            continue

        # 一级标题
        if line.startswith("# "):
            document.add_heading(
                line[2:].strip(),
                level=1
            )

        # 二级标题
        elif line.startswith("## "):
            document.add_heading(
                line[3:].strip(),
                level=2
            )

        # 三级标题
        elif line.startswith("### "):
            document.add_heading(
                line[4:].strip(),
                level=3
            )

        # 无序列表
        elif line.startswith("- "):
            document.add_paragraph(
                line[2:].strip(),
                style="List Bullet"
            )

        # 有序列表
        elif (
            len(line) >= 3
            and line[0].isdigit()
            and line[1] == "."
            and line[2] == " "
        ):
            document.add_paragraph(
                line[3:].strip(),
                style="List Number"
            )

        # 普通文本
        else:
            document.add_paragraph(line)

        i += 1


def create_word_report(
    title: str,
    content: str,
    output_path: str
):
    document = Document()

    # 报告标题
    document.add_heading(
        title,
        level=0
    )

    # Markdown 内容
    add_markdown_content(
        document,
        content
    )

    # 保存
    document.save(output_path)


def add_markdown_table(
    document: Document,
    lines: list[str]
):
    """
    将 Markdown 表格转换成 Word 表格。
    """

    rows = []

    for line in lines:
        line = line.strip()

        if not line.startswith("|"):
            continue

        cells = [
            cell.strip()
            for cell in line.strip("|").split("|")
        ]

        rows.append(cells)

    if len(rows) < 2:
        return

    # 第二行通常是 Markdown 分隔线
    if all(
        set(cell) <= {"-", ":"}
        for cell in rows[1]
    ):
        rows.pop(1)

    if not rows:
        return

    table = document.add_table(
        rows=1,
        cols=len(rows[0])
    )

    table.style = "Table Grid"

    # 表头
    header_cells = table.rows[0].cells

    for i, value in enumerate(rows[0]):
        header_cells[i].text = value

    # 数据行
    for row in rows[1:]:
        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = value